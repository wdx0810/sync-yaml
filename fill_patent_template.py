"""
严格按照「专利技术交底书模板（软通）.docx」的结构和格式填充内容。
规则：保留所有【】提示文字，将"例如：..."示例段落替换为实际内容。
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

TEMPLATE = r"C:\AI\project-configmap\project-configmap\专利技术交底书模板（软通）.docx"
OUTPUT = r"C:\AI\project-configmap\project-configmap\专利技术交底书-YAML同步系统.docx"
FIG_DIR = r"C:\AI\project-configmap\project-configmap\patent_figures"

doc = Document(TEMPLATE)

# ============================================================
# Helper: set paragraph text preserving the paragraph object
# ============================================================
def set_para(idx, text, bold=False):
    p = doc.paragraphs[idx]
    p.clear()
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(12)

# ============================================================
# 1. 填写表格元信息
# ============================================================
table = doc.tables[0]
for ci in range(1, 4):
    table.rows[0].cells[ci].paragraphs[0].clear()
    table.rows[0].cells[ci].paragraphs[0].add_run(
        "一种基于语义比较的GitLab与Kubernetes双向YAML资源同步方法及系统")
for ci in range(1, 4):
    table.rows[1].cells[ci].paragraphs[0].clear()
    table.rows[1].cells[ci].paragraphs[0].add_run("发明")

# ============================================================
# 2. 填写各章节内容（保留【】提示，替换"例如"示例）
# ============================================================

# --- 一、技术领域 ---
# P9 = 【提示】 保留
# P10 = "例如：..." → 替换为实际内容
set_para(10, "本申请涉及云原生容器编排与持续交付技术领域，特别涉及一种基于语义比较的GitLab与Kubernetes双向YAML资源同步方法及系统。")

# --- 二、背景技术 ---
# A1: P14=【提示】保留, P15="例如"→替换
set_para(15, "随着微服务架构和Kubernetes容器编排平台的广泛应用，企业通常将应用的声明式配置（Deployment、Service、ConfigMap等）以YAML文件形式存储在GitLab等代码托管平台中，通过GitOps流程管理集群资源。然而，运维团队经常需要在Kubernetes集群中直接修改资源（如紧急扩容、临时调整配置），导致集群运行态与GitLab存储态产生漂移（Configuration Drift）。目前业界主流GitOps工具主要实现单向同步，缺乏可靠的双向同步能力和细粒度审核机制。")

# A2: P18=【提示】保留, P19="例如"→替换
set_para(19, """现有最接近的技术方案为：
1）ArgoCD（开源项目，GitHub: argoproj/argo-cd）：通过持续监控Git仓库并与集群状态对比，实现Git→K8s的单向声明式同步，支持自动或手动同步模式；
2）FluxCD（开源项目，GitHub: fluxcd/flux2）：通过Source Controller监听Git仓库变化，Kustomize Controller自动应用到集群，同样聚焦Git→K8s单向流程；
3）专利CN115086177A"一种基于GitOps的Kubernetes资源同步方法"公开了通过Webhook触发Git到K8s的单向同步方案，但未涉及K8s到Git的反向回写和语义比较机制。""")

# A3: P22=【提示】保留, P23="例如"→替换
set_para(23, """现有技术存在以下缺陷和不足：
1）仅支持单向同步（Git→K8s），无法将Kubernetes集群中的运行时变更自动回写到Git仓库，导致Git仓库与集群实际状态持续漂移，违背了"Git作为唯一事实源"的GitOps原则；
2）反向同步时，Kubernetes API Server会自动填充大量默认值字段（如Deployment的revisionHistoryLimit:10、progressDeadlineSeconds:600、strategy.rollingUpdate默认值，Pod模板的terminationGracePeriodSeconds:30、dnsPolicy:ClusterFirst、restartPolicy:Always、容器的imagePullPolicy:IfNotPresent等），以及运行时字段（status、resourceVersion、uid、managedFields、creationTimestamp等）。直接序列化会将这些字段写入Git，导致Git中的YAML与用户编写的原始YAML严重不一致，再次正向同步时产生大量无意义的变更甚至"field is immutable"错误；
3）现有工具的正向同步缺乏变更审核机制，直接将Git中的YAML全量应用到集群，可能造成生产事故；
4）由于数字类型在不同解析器中表示不一致（sigs.k8s.io/yaml通过JSON路由产生float64，而Kubernetes API返回int64），使用简单的字符串比较或reflect.DeepEqual判断资源是否变化会产生大量误报（如int64(80)≠float64(80)被判为有变更），导致不必要的同步操作和Git提交；
5）大规模资源（数百到上千个）同步时，逐个API调用导致同步耗时过长（测试环境930个资源需85秒），无法满足生产环境的实时性要求；
6）缺乏面向开发人员的受控配置修改入口：开发人员修改配置需要直接操作Git仓库或集群，缺乏"编辑—审核—提交"的闭环流程；且在多人同时修改同一配置文件时，缺乏并发冲突控制机制，后提交者可能无意中覆盖先提交者的改动，造成配置丢失。""")

# --- 三、发明内容 ---
# B1: P27=【提示】保留, P28="例如"→替换
set_para(28, """本申请要解决的技术问题包括：
1）解决GitLab与Kubernetes之间缺乏双向同步导致的配置漂移问题；
2）解决反向同步中Kubernetes默认值和运行时字段污染导致的往返不一致问题；
3）解决正向同步缺乏审核机制可能导致的生产安全问题；
4）解决不同数据来源的数字类型不一致导致语义相同资源被误判为变更的问题；
5）解决大规模资源同步时串行处理导致效率低下的问题。""")

# B2: P31=【提示】保留, P32="例如"→替换, P33~P36也是例子→清空
set_para(32, """本申请提出一种基于语义比较的GitLab与Kubernetes双向YAML资源同步方法，包括以下技术方案：

方案一：运行时字段清洗机制（Cleaner）
设计一个多层次的字段清洗器，对从Kubernetes API获取的资源对象执行以下清洗操作：
（1）移除通用运行时字段：status、managedFields、resourceVersion、uid、creationTimestamp、generation、selfLink等；
（2）移除Kubernetes自动注入的注解：kubectl.kubernetes.io/last-applied-configuration、deployment.kubernetes.io/revision及kubernetes.io/*前缀注解；
（3）按资源类型（Kind）移除API Server填充的默认值：
  - Service：移除clusterIP、clusterIPs、ipFamilies、ipFamilyPolicy、internalTrafficPolicy等，若type为ClusterIP（默认值）则移除，端口中protocol为TCP（默认值）则移除，targetPort与port相同则移除；
  - Deployment/StatefulSet/DaemonSet：若revisionHistoryLimit=10则移除、若progressDeadlineSeconds=600则移除、若strategy为默认RollingUpdate(maxSurge:25%,maxUnavailable:25%)则移除；
  - Pod模板：restartPolicy=Always则移除、terminationGracePeriodSeconds=30则移除、dnsPolicy=ClusterFirst则移除、schedulerName=default-scheduler则移除、空securityContext则移除；
  - 容器：imagePullPolicy=IfNotPresent则移除、terminationMessagePath=/dev/termination-log则移除、terminationMessagePolicy=File则移除、端口protocol=TCP则移除、探针successThreshold=1/failureThreshold=3/periodSeconds=10/timeoutSeconds=1则移除；
（4）移除ServiceAccount自动注入的Volume和VolumeMount（kube-api-access-*、default-token-*前缀）。

方案二：JSON归一化语义比较（IsSameContent）
设计统一的语义比较方法，正向同步和反向同步均使用完全相同的比较逻辑：
（1）对两个待比较的资源对象分别执行Clean清洗；
（2）提取可比较字段（排除apiVersion、kind、metadata中的非语义字段，保留spec、data、metadata.labels、metadata.annotations）；
（3）对提取的字段递归执行数值归一化（normalize）：int/int32/int64统一转为float64，移除nil值、空map(len==0)和空slice(len==0)；
（4）将归一化后的对象序列化为JSON字符串进行比较。

方案三：带审核的正向同步流程
（1）预览阶段：批量List获取集群现有资源建立内存索引，与GitLab解析的YAML逐项执行IsSameContent语义比较，生成变更列表含YAML差异；
（2）审核阶段：用户在Web界面逐项审查diff并勾选需要应用的变更；
（3）应用阶段：使用Server-Side Apply并发应用已批准的变更到集群。

方案四：原子批量反向同步
（1）并发List所有目标GVR+命名空间组合的资源；
（2）对每个资源执行Clean清洗后序列化为YAML；
（3）将GitLab中已有文件通过Parse解析为对象，使用IsSameContent执行语义比较（而非字符串比较）；
（4）仅对语义确有变化的资源收集变更，通过单次原子Commit批量写入GitLab。

方案五：多级并发处理架构
（1）GitLab文件拉取使用20个并发Worker；
（2）Kubernetes资源List使用10个并发Worker；
（3）变更Apply使用10个并发Worker；
（4）Watch模式下设置5秒防抖窗口，单次编辑只产生一个commit。

方案六：基于乐观锁的受控配置修改与并发冲突控制
设计一套"加载—编辑—提交—审核—写入"的受控配置修改流程，并通过乐观锁解决多人并发修改冲突：
（1）加载阶段：用户从GitLab加载目标配置文件的当前内容，系统同时计算该内容的哈希值作为"版本基线"随修改申请一并保存；编辑界面默认只读，需显式进入编辑态并保存后方可提交；
（2）提交阶段：若同一文件已存在待审核的修改申请，系统提示申请人存在并发修改风险，由申请人确认后再创建申请；
（3）审核写入阶段：审核人批准时，系统重新读取GitLab中该文件的当前内容并计算哈希，与申请保存的版本基线比对：若一致则提交写入GitLab；若不一致（说明该文件在此期间已被其他申请合入），则拒绝本次提交，将该申请自动标记为"已失效"，并保存当前最新内容用于差异展示；
（4）冲突呈现阶段：对已失效的申请，系统向用户展示"版本基线与GitLab最新内容"以及"版本基线与本申请修改内容"两组差异，供申请人对照后基于最新内容重新编辑提交，实现人工可控的冲突合并。""")
for i in range(33, 37):
    set_para(i, "")

# B3: P39=【提示】保留, P40="例如"→替换
set_para(40, """本申请与现有技术相比具有以下有益效果：
1）通过运行时字段清洗机制和JSON归一化语义比较，实现了"往返一致性"——即从Kubernetes读出经清洗后写入GitLab的YAML，再正向同步回Kubernetes时不会产生虚假变更，解决了现有技术中默认值污染和类型不匹配导致的同步循环问题；
2）正反向同步使用完全相同的比较逻辑（IsSameContent），消除了因比较方法不一致导致的"正向同步0变更但反向同步报大量变更"的矛盾；
3）带审核的正向同步流程，使运维人员能在变更应用到生产集群前逐项审查YAML差异，降低了误操作风险；
4）多级并发处理使930+资源的完整同步从85秒降低到33秒（提升约61%），满足生产环境实时性要求；
5）反向同步的单次原子Commit避免了一次K8s变更产生多个Git提交的问题，保持Git历史清晰可追溯；
6）通过受控配置修改流程和乐观锁并发控制，为开发人员提供了"编辑—审核—提交"的闭环入口，且在多人并发修改同一配置时能可靠检测冲突、防止覆盖，相比直接操作Git仓库或引入实时协同编辑框架，既保证了安全性又避免了系统复杂度的大幅上升。""")

# B4: P43=【提示】保留, P44="例如"→替换
set_para(44, """本申请的主要创新点在于：
1）提出了一种面向Kubernetes资源的多层次运行时字段清洗方法，能按资源Kind精确识别并移除API Server填充的默认值字段，实现Clean(K8sLive)≈Clean(GitLabYAML)的往返稳定性；
2）提出了一种基于JSON归一化的语义比较方法，通过将所有数值类型统一转换为float64并移除空容器，消除了不同YAML/JSON解析器产生的类型差异；
3）提出了正向和反向同步必须使用完全相同比较逻辑的一致性约束，从架构层面消除了双向同步中的"假变更"问题；
4）设计了"预览→审核→应用"三阶段正向同步流程，在保证GitOps自动化的同时引入人工审核环节；
5）采用批量List建立内存索引替代逐个Get的方式进行资源比较，配合多级并发Worker池，实现大规模资源集的高效同步；
6）提出了一种基于内容哈希版本基线的乐观锁并发控制方法，应用于"加载—编辑—审核—提交"的受控配置修改流程，在多人同时修改同一配置文件时，能在提交写入前检测版本冲突、拒绝覆盖，并展示双向差异供人工合并，从而在不引入实时协同编辑复杂度的前提下保证配置不被误覆盖。""")

# --- 四、专利附图 ---
# P47=【提示】保留, P48="例如"→替换, P49~P54→清空后插图
set_para(48, """图1为本发明系统整体架构示意图；
图2为本发明运行时字段清洗（Cleaner）流程图；
图3为本发明JSON归一化语义比较（IsSameContent）流程图；
图4为本发明正向同步（GitLab→K8s）带审核流程图；
图5为本发明反向同步（K8s→GitLab）原子批量提交流程图；
图6为本发明多级并发处理架构图。""")

# P49~P54: clear examples, then insert figures
for i in range(49, 55):
    set_para(i, "")

# Insert figures into P49~P54
for fig_num in range(1, 7):
    fig_path = os.path.join(FIG_DIR, f"fig{fig_num}.png")
    target_idx = 48 + fig_num  # P49, P50, P51, P52, P53, P54
    if os.path.exists(fig_path) and target_idx < len(doc.paragraphs):
        p = doc.paragraphs[target_idx]
        p.clear()
        run = p.add_run(f"图{fig_num}")
        run.font.size = Pt(10)
        run = p.add_run()
        run.add_picture(fig_path, width=Cm(14))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# --- 五、具体实施方式 ---
# P56=【提示】保留, P57="例如"→替换, P58~P66也是例子→部分清空
set_para(57, """本发明实施例提供的方法结合附图详细描述如下：

实施例一：运行时字段清洗（参见图2）

当系统从Kubernetes API获取到一个Deployment资源对象时，清洗器（Cleaner）执行以下步骤：
S1，深拷贝原始对象，确保不修改Kubernetes客户端缓存中的数据；
S2，移除顶层status字段（该字段为纯运行时状态，不可声明式应用）；
S3，遍历metadata字段，移除managedFields、resourceVersion、uid、creationTimestamp、generation、selfLink、deletionTimestamp、ownerReferences、finalizers等运行时字段；
S4，遍历metadata.annotations，移除kubectl.kubernetes.io/last-applied-configuration、deployment.kubernetes.io/revision等系统注解，以及所有kubernetes.io/*、k8s.io/*前缀的注解；若移除后annotations为空则删除整个annotations字段；
S5，识别资源Kind为"Deployment"，执行特定默认值清洗：
  - 若spec.revisionHistoryLimit等于10（API Server默认值），则移除该字段；
  - 若spec.progressDeadlineSeconds等于600（API Server默认值），则移除该字段；
  - 若spec.strategy.type为"RollingUpdate"且rollingUpdate.maxSurge为"25%"、maxUnavailable为"25%"（均为默认值），则移除rollingUpdate子字段；若strategy中仅剩type一个字段且为默认值，则移除整个strategy字段；
S6，进入Pod模板（spec.template）清洗：
  - 移除spec.template.metadata.creationTimestamp（client-go序列化的噪音字段）；
  - 移除冗余的spec.template.spec.serviceAccount字段（保留serviceAccountName即可）；
  - 若restartPolicy为"Always"（默认值）则移除；
  - 若terminationGracePeriodSeconds为30（默认值）则移除；
  - 若dnsPolicy为"ClusterFirst"（默认值）则移除；
  - 若schedulerName为"default-scheduler"（默认值）则移除；
  - 若securityContext为空对象{}则移除；
  - 过滤自动注入的ServiceAccount Token卷（名称前缀为kube-api-access-或default-token-）；
S7，对每个容器执行清洗：
  - 若imagePullPolicy为"IfNotPresent"（非:latest标签的默认值）则移除；
  - 若terminationMessagePath为"/dev/termination-log"（默认值）则移除；
  - 若terminationMessagePolicy为"File"（默认值）则移除；
  - 移除端口定义中protocol为"TCP"（默认值）的字段；
  - 过滤自动注入的volumeMounts（名称前缀为kube-api-access-或default-token-）；
  - 对探针（livenessProbe/readinessProbe/startupProbe）移除等于默认值的字段：successThreshold=1、failureThreshold=3、periodSeconds=10、timeoutSeconds=1。

清洗后的对象仅保留用户显式指定的字段，序列化为YAML后可在任意集群中直接kubectl apply。

实施例二：JSON归一化语义比较（参见图3）

当需要判断两个资源对象（一个来自Kubernetes Live，一个来自GitLab文件）是否语义相同时，执行以下步骤：
S1，分别对两个对象执行Clean清洗（同实施例一）；
S2，从清洗后的对象中提取可比较字段：遍历对象顶层key，排除apiVersion、kind、metadata中已清洗的非业务字段，保留spec、data等业务字段，同时包含metadata.labels和metadata.annotations；
S3，对提取的字段执行递归归一化（normalize函数）：
  - 遇到map[string]interface{}类型：递归处理每个value，若value归一化后为nil、空map或空slice则跳过（不放入结果）；
  - 遇到[]interface{}类型：递归处理每个元素；
  - 遇到int、int32、int64类型的值：转换为float64（因为Kubernetes API返回int64，而sigs.k8s.io/yaml解析器通过JSON中间层产生float64，reflect.DeepEqual会将int64(80)与float64(80)判定为不相等）；
  - 其他类型（string、bool、float64）：保持不变；
S4，将归一化后的两个map分别通过json.Marshal序列化为JSON字节串；
S5，比较两个JSON字符串是否完全相等：相等则判定资源语义相同，返回true跳过同步；不相等则判定有实际变更，需要同步。

关键设计约束：正向同步和反向同步中判断"是否有变更"的逻辑必须使用完全相同的IsSameContent方法。若一侧使用语义比较、另一侧使用字符串比较，会导致"正向同步判定0变更但反向同步报大量变更"的往返不一致问题。

实施例三：带审核的正向同步（参见图4）

正向同步（GitLab→K8s）执行以下步骤：
S1（预览阶段 - 拉取GitLab文件），系统使用20个并发Worker从GitLab API拉取指定路径下的所有YAML文件内容；
S2（预览阶段 - 解析资源），使用sigs.k8s.io/yaml解析每个YAML文件为Kubernetes Unstructured对象，通过GVR Resolver解析apiVersion+kind得到GroupVersionResource和是否为命名空间级资源；
S3（预览阶段 - 批量List），确定所有需要比较的GVR+命名空间唯一组合（通常约5~64种），使用10个并发Worker执行List操作获取集群中已有资源，建立内存索引（key为"GVR字符串|命名空间|资源名"→Unstructured对象）；
S4（预览阶段 - 语义比较），遍历从GitLab解析出的每个资源，在内存索引中O(1)查找对应的集群资源：
  - 若找到：执行IsSameContent比较，相同则标记"跳过"，不同则标记"待更新"并生成Clean后的YAML diff供用户审查；
  - 若未找到：标记"待创建"；
S5（审核阶段），前端展示所有"待更新"和"待创建"资源列表，总数/变更数/无变更数统计，每项以并排diff形式显示（左：集群当前Clean后的YAML，右：GitLab期望Clean后的YAML），用户可逐项勾选或取消；
S6（应用阶段），用户点击"应用变更"后，系统对已勾选的变更项使用10个并发Worker执行Server-Side Apply（通过Dynamic Client的Patch方法，patchType为ApplyPatch），返回每项的成功/失败结果及错误详情。

实施例四：原子批量反向同步（参见图5）

反向同步（K8s→GitLab）执行以下步骤：
S1，根据任务配置的资源类型列表和命名空间列表，构建所有GVR+命名空间的List任务（如任务配置了ConfigMap/Deployment/Service三种类型和两个命名空间，则产生6个List任务）；
S2，使用10个并发Worker执行所有List任务获取集群资源列表；
S3，对每个获取到的资源依次执行过滤：
  - 有OwnerReferences的资源（如Deployment拥有的ReplicaSet）跳过；
  - 系统资源跳过（kube-root-ca.crt ConfigMap、default ServiceAccount、kubernetes.io/service-account-token类型Secret、system:前缀的ClusterRole/ClusterRoleBinding等）；
  - 执行用户配置的includeFilter/excludeFilter正则匹配；
S4，对通过过滤的资源执行Clean清洗，使用parser.Print序列化为YAML字节串；
S5，计算该资源在GitLab中的存储路径（格式：{basePath}/{namespace}/{resourceTypePlural}/{name}.yaml）；
S6，在预先拉取的GitLab现有文件内容中查找该路径：
  - 若存在：将现有文件内容通过parser.Parse解析为Unstructured对象，然后执行IsSameContent语义比较（不使用字符串比较）。相同则跳过，不同则加入变更列表；
  - 若不存在：标记为新建文件，加入变更列表；
S7，将所有变更文件通过GitLab Commits API打包为单次原子Commit提交（CreateCommit，包含多个CommitAction），commit message格式为"[任务名称] Sync from K8s: N resource(s)"。

实施例五：受控配置修改与乐观锁并发冲突控制

本实施例描述开发人员在线修改ConfigMap配置并经审核提交至GitLab的完整流程，重点说明多人并发场景下的冲突控制：
S1，开发人员甲在配置修改界面选择目标环境（对应某GitLab仓库路径）和目标ConfigMap，系统从GitLab读取该文件当前内容C1，计算其内容哈希H1=hash(C1)作为版本基线；编辑框默认为只读状态；
S2，甲点击"编辑"进入可编辑状态，修改内容后点击"保存修改"暂存为待提交内容C1'，界面回到只读态并展示C1与C1'的差异预览；
S3，甲填写变更说明并点击"提交审核"，系统检查是否存在针对同一文件的其他待审核申请：若存在，弹出确认对话框告知甲存在并发修改（含待审数量与申请人），甲确认后创建修改申请R1（保存版本基线H1、原内容C1、修改内容C1'、变更说明、申请人）；
S4，与此同时，开发人员乙同样基于C1（版本基线H1）修改该文件为C1''并提交，创建申请R2；此时R1、R2均处于待审核状态且基线相同；
S5，审核人先批准R1：系统重新读取GitLab中该文件的当前内容，计算哈希并与R1的基线H1比对，一致，于是将C1'提交写入GitLab（commit message包含申请人甲、审核人信息），GitLab该文件更新为C1'，其哈希变为H2；
S6，审核人再批准R2：系统重新读取GitLab当前内容（此时为C1'，哈希H2），与R2的基线H1比对，不一致，判定为版本冲突；系统拒绝提交，将R2状态自动置为"已失效(冲突)"，并将当前最新内容C1'保存至R2的冲突内容字段；
S7，乙查看R2详情时，界面展示两组并排差异：①版本基线C1 与 GitLab最新C1'（即甲合入的改动），②版本基线C1 与 乙的修改C1''（即乙原本的改动）；乙据此对照，基于最新内容C1'重新编辑并提交新申请，完成人工可控的冲突合并。

通过上述流程，系统在不引入CRDT/OT等实时协同编辑框架的前提下，利用内容哈希版本基线的乐观锁机制，可靠地防止了后提交者覆盖先提交者改动的问题。""")

# Clear remaining example paragraphs in section 5
for i in range(58, 67):
    if i < len(doc.paragraphs):
        set_para(i, "")

# --- 六、替代方案 ---
# P68=【提示】保留, after it add content
alt_idx = None
for i, p in enumerate(doc.paragraphs):
    if "六、替代方案" in p.text:
        alt_idx = i
        break

if alt_idx:
    # The instruction paragraph is alt_idx+1 (P68), keep it
    # The paragraph after instruction (alt_idx+2) is where we write
    target = alt_idx + 2
    if target < len(doc.paragraphs):
        set_para(target, """针对本申请的技术方案，存在以下替代方案：

1）运行时字段清洗的替代方案：可以使用白名单模式（仅保留用户显式指定的字段）替代本申请的黑名单模式（移除已知的默认值字段）。白名单模式通过维护一个"用户可能设置的字段列表"来决定保留哪些字段，但由于Kubernetes资源Kind众多且字段不断演进，黑名单模式更易维护且不会误删用户有意设置的非默认值字段；

2）语义比较的替代方案：可以使用结构化diff算法（如递归逐字段比较并记录差异路径）替代JSON序列化比较。两种方案的比较结果相同，但JSON序列化方式实现更简洁，且天然解决了Go语言map迭代顺序不确定的问题；

3）审核流程的替代方案：可以使用Pull Request/Merge Request模式——系统将待应用的变更写入GitLab的临时分支并创建Merge Request，人工审核合并后再触发正向同步。该方案审核记录更完整（利用GitLab原生的Review机制），但增加了操作步骤和延迟；

4）并发架构的替代方案：可以使用消息队列（如NATS、Redis Stream、RabbitMQ）替代Go语言channel+WaitGroup的进程内并发模型，适合多实例分布式部署场景；对于单实例部署，进程内并发模型性能更优且运维更简单；

5）反向同步触发的替代方案：除Watch模式（实时监听K8s事件）外，也可使用定时轮询模式（如每5分钟执行一次List+比较）。Watch模式实时性更好，但长时间运行可能因连接断开需要重连重试机制；定时轮询模式实现更简单但有延迟；

6）并发冲突控制的替代方案：本申请采用基于内容哈希的乐观锁（提交时校验）。替代方案一为悲观锁——用户开始编辑时即锁定该文件，其他人无法编辑，直到提交或超时释放；但悲观锁在用户长时间占用或异常退出时会阻塞他人，体验较差。替代方案二为实时协同编辑（引入CRDT或OT算法配合WebSocket长连接），支持多人同时编辑同一文档；但其实现复杂度高，且对配置文件这类不需要多人同时录入的场景属于过度设计。相比之下，本申请的乐观锁方案实现简单、无锁等待、且能可靠防止覆盖，最契合配置修改审核场景。""")

# --- 七、其他事项 ---
# Keep as-is (template placeholders for user to fill manually)

# ============================================================
# Save
# ============================================================
doc.save(OUTPUT)
print(f"Done! Saved to: {OUTPUT}")
